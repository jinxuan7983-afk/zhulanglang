
from docx import Document
import os

def read_word_document(file_path):
    doc = Document(file_path)
    content = []
    
    for para in doc.paragraphs:
        if para.text.strip():
            content.append(para.text)
    
    for table in doc.tables:
        content.append('\n[表格]\n')
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells]
            if any(row_data):
                content.append(' | '.join(row_data))
    
    return '\n'.join(content)

docs = [
    '需求规格说明书_2301130119_朱郎郎.docx',
    '2.+系统界面设计说明书_学号_姓名_v1.docx',
    '3. 数据库设计说明书_2301130119_朱郎郎小组_v1.docx'
]

for doc_file in docs:
    if os.path.exists(doc_file):
        output_file = doc_file.replace('.docx', '.txt')
        content = read_word_document(doc_file)
        with open(output_file, 'w', encoding='utf-8') as f:
            f.write(content)
        print(f'已读取: {doc_file} -> {output_file}')
